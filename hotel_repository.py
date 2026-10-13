"""SQLite hotel/offer filtering and FAISS retrieval with pluggable embeddings."""
from __future__ import annotations

import hashlib
import json
import sqlite3
import threading
import time
from datetime import date, timedelta
from pathlib import Path
from contextlib import contextmanager

import faiss
import numpy as np


class DemoEncoder:
    """Offline character-bigram hash embeddings, explicitly not BGE-M3."""
    name = "demo-hash-384"

    def encode(self, texts):
        vectors = np.zeros((len(texts), 384), dtype="float32")
        for row, text in enumerate(texts):
            tokens = list(text) + [text[i:i + 2] for i in range(len(text) - 1)]
            for token in tokens:
                digest = hashlib.blake2b(token.encode(), digest_size=4).digest()
                vectors[row, int.from_bytes(digest, "little") % 384] += 1
        faiss.normalize_L2(vectors)
        return vectors


class BGEEncoder:
    def __init__(self, model_path="BAAI/bge-m3"):
        from sentence_transformers import SentenceTransformer
        self.name = "bge-m3:" + model_path
        self.model = SentenceTransformer(model_path)

    def encode(self, texts):
        return np.asarray(self.model.encode(texts, normalize_embeddings=True,
                                           show_progress_bar=False), dtype="float32")


class HotelRepository:
    def __init__(self, db_path: str, encoder=None):
        self.db_path = str(db_path)
        Path(self.db_path).parent.mkdir(parents=True, exist_ok=True)
        self.encoder = encoder or DemoEncoder()
        self.index = None
        self.index_ids = []
        self.index_lock = threading.Lock()
        with self.connect() as db:
            db.executescript("""
                PRAGMA journal_mode=WAL;
                CREATE TABLE IF NOT EXISTS hotels (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, city TEXT NOT NULL,
                    city_id TEXT, brand TEXT);
                CREATE INDEX IF NOT EXISTS idx_hotels_city ON hotels(city);
                CREATE TABLE IF NOT EXISTS offers (
                    hotel_id TEXT NOT NULL, night TEXT NOT NULL, room_type TEXT NOT NULL,
                    price REAL NOT NULL CHECK(price>0), remaining INTEGER NOT NULL CHECK(remaining>=0),
                    source TEXT NOT NULL, updated TEXT NOT NULL,
                    PRIMARY KEY(hotel_id,night,room_type));
                CREATE INDEX IF NOT EXISTS idx_offers_filter ON offers(night,room_type,price);
            """)
            columns = {row[1] for row in db.execute("PRAGMA table_info(offers)")}
            for column in ("observed_at", "valid_until"):
                if column not in columns:
                    db.execute(f"ALTER TABLE offers ADD COLUMN {column} REAL")

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.db_path, timeout=10)
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def import_catalog(self, xlsx_path: str):
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_path, read_only=True)
        rows = []
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                rows.append(tuple(str(value or "") for value in (row[1], row[0], row[2], row[3], row[4])))
        workbook.close()
        with self.connect() as db:
            db.executemany("INSERT INTO hotels VALUES(?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,city=excluded.city,city_id=excluded.city_id,brand=excluded.brand", rows)
        self.index = None
        return self.count()

    def count(self):
        with self.connect() as db:
            return db.execute("SELECT COUNT(*) FROM hotels").fetchone()[0]

    def cities(self):
        with self.connect() as db:
            return [row[0] for row in db.execute("SELECT DISTINCT city FROM hotels ORDER BY city")]

    def catalog(self):
        with self.connect() as db:
            return [dict(row) for row in db.execute("SELECT * FROM hotels ORDER BY id")]

    def put_offer(self, hotel_id, night, room_type, price, remaining, source="live", ttl=900):
        with self.connect() as db:
            assert db.execute("SELECT id FROM hotels WHERE id=?", (hotel_id,)).fetchone(), "Unknown hotel ID"
            if ttl <= 0:
                raise ValueError("Offer TTL must be positive")
            db.execute("INSERT INTO offers(hotel_id,night,room_type,price,remaining,source,updated,observed_at,valid_until) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(hotel_id,night,room_type) DO UPDATE SET price=excluded.price,remaining=excluded.remaining,source=excluded.source,updated=excluded.updated,observed_at=excluded.observed_at,valid_until=excluded.valid_until",
                       (hotel_id, night, room_type, price, remaining, source, date.today().isoformat(), time.time(), time.time() + ttl))

    def build_index(self):
        hotels = self.catalog()
        if not hotels:
            raise ValueError("Hotel catalog is empty")
        self.index_ids = [hotel["id"] for hotel in hotels]
        fingerprint = hashlib.sha256(json.dumps(
            [self.encoder.name, [(h["id"], h["name"], h["city"]) for h in hotels]],
            ensure_ascii=False).encode()).hexdigest()
        index_path = Path(self.db_path).with_suffix(".faiss")
        meta_path = Path(self.db_path).with_suffix(".index.json")
        if index_path.exists() and meta_path.exists():
            try:
                metadata = json.loads(meta_path.read_text(encoding="utf-8"))
                if metadata.get("fingerprint") == fingerprint:
                    cached = faiss.read_index(str(index_path))
                    if cached.ntotal == len(self.index_ids):
                        self.index = cached
                        return
            except (ValueError, RuntimeError, OSError):
                pass  # Rebuild corrupt or stale cache from the durable catalog.
        vectors = self.encoder.encode([hotel["city"] + " " + hotel["name"] for hotel in hotels])
        self.index = faiss.IndexFlatIP(vectors.shape[1])
        self.index.add(vectors)
        temporary = str(index_path) + ".tmp"
        faiss.write_index(self.index, temporary)
        Path(temporary).replace(index_path)
        meta_path.write_text(json.dumps({"fingerprint": fingerprint, "encoder": self.encoder.name}), encoding="utf-8")

    def search(self, query: str, slots: dict, limit=5):
        if self.index is None:
            with self.index_lock:
                if self.index is None:
                    self.build_index()
        # Rank the whole catalog; SQLite applies the hard business constraints.
        # This avoids falsely returning no matches when top-k semantic recall misses an eligible offer.
        scores, positions = self.index.search(self.encoder.encode([query]), self.index.ntotal)
        ranks = {self.index_ids[int(pos)]: float(score) for score, pos in zip(scores[0], positions[0]) if pos >= 0}
        start, end = date.fromisoformat(slots["check_in"]), date.fromisoformat(slots["check_out"])
        nights = (end - start).days
        with self.connect() as db:
            # Every night must exist, be available, and satisfy the per-night budget.
            rows = db.execute("""
                SELECT h.id,h.name,h.city,h.brand,o.room_type,
                       SUM(o.price) total_price,MAX(o.price) max_nightly_price,
                       MIN(o.remaining) remaining,GROUP_CONCAT(DISTINCT o.source) sources,
                       MIN(o.valid_until) valid_until,MAX(o.observed_at) observed_at
                FROM hotels h JOIN offers o ON h.id=o.hotel_id
                WHERE h.city=? AND o.night>=? AND o.night<? AND o.room_type=? AND o.valid_until>?
                GROUP BY h.id,o.room_type
                HAVING COUNT(DISTINCT o.night)=? AND MIN(o.remaining)>0
                       AND (CASE WHEN ?='stay' THEN SUM(o.price) ELSE MAX(o.price) END)<=?
            """, (slots["city"], start.isoformat(), end.isoformat(), slots["room_type"], time.time(), nights, slots.get("budget_scope", "night"), slots["budget"])).fetchall()
            city_offers = db.execute("SELECT COUNT(*) FROM offers o JOIN hotels h ON o.hotel_id=h.id WHERE h.city=? AND o.night>=? AND o.night<? AND o.valid_until>?",
                                     (slots["city"], start.isoformat(), end.isoformat(), time.time())).fetchone()[0]
            hotel_count = db.execute("SELECT COUNT(*) FROM hotels WHERE city=?", (slots["city"],)).fetchone()[0]
        candidates = [{**dict(row), "score": ranks.get(row["id"], 0), "nights": nights} for row in rows]
        for hotel in candidates:
            signature = [hotel["id"], slots["check_in"], slots["check_out"], hotel["room_type"],
                         hotel["total_price"], hotel["valid_until"]]
            hotel["quote_id"] = hashlib.sha256(json.dumps(signature).encode()).hexdigest()[:32]
        candidates.sort(key=lambda hotel: (-hotel["score"], hotel["total_price"], hotel["id"]))
        return {"candidates": candidates[:limit], "eligible_count": len(candidates),
                "city_hotels": hotel_count, "offers_known": city_offers > 0,
                "encoder": self.encoder.name, "catalog_count": self.count()}

    def suggest(self, query, city, limit=2):
        if self.index is None:
            with self.index_lock:
                if self.index is None:
                    self.build_index()
        scores, positions = self.index.search(self.encoder.encode([query]), self.index.ntotal)
        hotels = {hotel["id"]: hotel for hotel in self.catalog() if hotel["city"] == city}
        result = []
        for position, score in zip(positions[0], scores[0]):
            hotel = hotels.get(self.index_ids[int(position)])
            if hotel:
                result.append({**hotel, "score": float(score)})
                if len(result) >= limit:
                    break
        return result

    def seed_demo_offers(self, start: date, days=7):
        """Use real catalog names with explicitly synthetic, isolated sample inventory."""
        hotels = self.catalog()
        selected = []
        for city in ("上海", "北京", "杭州", "成都"):
            selected.extend([hotel for hotel in hotels if hotel["city"] == city][:4])
        for offset in range(days):
            night = (start + timedelta(days=offset)).isoformat()
            for i, hotel in enumerate(selected):
                for room_type, premium in (("大床房", 0), ("双床房", 30)):
                    self.put_offer(hotel["id"], night, room_type, 220 + (i % 4) * 45 + premium, 3, "demo_fixture")
        return selected
