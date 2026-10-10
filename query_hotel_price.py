"""
酒店每日房价查询工具 — 完全独立运行

用法: python query_hotel_price.py <距今天数1> <距今天数2> <酒店名>

示例: python query_hotel_price.py 10 12 全季酒店
  → 从今天起第10天入住、第12天离店，逐日显示每晚房价
"""
import asyncio
import difflib
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional

from loguru import logger
from openpyxl import load_workbook
from playwright.async_api import async_playwright


def _run_async(coro):
    with ThreadPoolExecutor() as pool:
        return pool.submit(asyncio.run, coro).result()


HOTELS_DB_PATH = os.path.join(os.path.dirname(__file__), "data", "hotels.xlsx")


def _normalize(name: str) -> str:
    name = re.sub(r'\([^)]*\)|（[^）]*）|\[[^\]]*\]|【[^】]*】', '', name)
    name = re.sub(r'[^一-龥a-zA-Z0-9]', '', name)
    cn_num_map = {'零': 0, '一': 1, '二': 2, '三': 3, '四': 4,
                  '五': 5, '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
    def replace_cn(match):
        cn = match.group(0)
        if len(cn) == 2 and cn[0] in cn_num_map and cn[1] == '十':
            return str(cn_num_map[cn[0]] * 10)
        elif len(cn) == 2 and cn[0] == '十' and cn[1] in cn_num_map:
            return str(10 + cn_num_map[cn[1]])
        elif len(cn) == 3 and cn[1] == '十':
            return str(cn_num_map[cn[0]] * 10 + cn_num_map[cn[2]])
        return str(cn_num_map.get(cn[0], 0))
    name = re.sub(r'[零一二三四五六七八九十]{1,3}', replace_cn, name)
    for word in ['酒店', '分店', '大厦店', '楼店', '层店', '新店', '店']:
        name = name.replace(word, '')
    return name.strip()


def _load_hotel_db(path: str = HOTELS_DB_PATH) -> List[Dict[str, str]]:
    if not os.path.exists(path):
        logger.warning(f"酒店ID数据库不存在: {path}")
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    if ws.max_row <= 1:
        wb.close()
        return []
    hotels = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        hotels.append({
            "酒店名": str(row[0] or ""),
            "酒店ID": str(row[1] or ""),
            "城市名": str(row[2] or ""),
            "城市ID": str(row[3] or ""),
            "品牌名": str(row[4] or ""),
            "品牌类型": str(row[5] or ""),
        })
    wb.close()
    return hotels


def _find_hotel_by_name(hotels: List[Dict], keyword: str, threshold: float = 0.5) -> Optional[Dict]:
    kw = _normalize(keyword)
    best = None
    best_score = 0.0
    for h in hotels:
        name = _normalize(h["酒店名"])
        if kw in name or name in kw:
            score = 0.9
        else:
            score = difflib.SequenceMatcher(None, kw, name).ratio()
        if score > best_score and score >= threshold:
            best_score = score
            best = h
    return best


def _query_night_price_sync(hotel_name: str, check_in: str, check_out: str):
    """查询单晚房价（同步接口，内部启动 Playwright）"""
    hotel_db = _load_hotel_db()
    if not hotel_db:
        return {"error": "酒店ID数据库为空，请确认 data/hotels.xlsx 存在"}

    matched = _find_hotel_by_name(hotel_db, hotel_name)
    if not matched:
        return {"error": f"未找到 '{hotel_name}' 对应的酒店记录"}

    hotel_id = matched["酒店ID"]
    matched_name = matched["酒店名"]
    logger.info(f"匹配酒店: {matched_name}, ID: {hotel_id}")

    detail_url = (
        f"https://hrewards.huazhu.com/hotel/detail?"
        f"checkInDate={check_in}&checkOutDate={check_out}&hotelId={hotel_id}"
    )

    async def _crawl():
        async with async_playwright() as p:
            browser = await p.chromium.launch(
                headless=True,
                args=[
                    "--disable-blink-features=AutomationControlled",
                    "--disable-dev-shm-usage",
                    "--no-sandbox",
                    "--disable-gpu",
                    "--disable-setuid-sandbox",
                    "--disable-web-security",
                ],
            )
            context = await browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                viewport={"width": 1920, "height": 1080},
                locale="zh-CN",
                timezone_id="Asia/Shanghai",
            )
            await context.add_init_script("""
                Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
                Object.defineProperty(navigator, 'plugins', {get: () => [1, 2, 3, 4, 5]});
                Object.defineProperty(navigator, 'languages', {get: () => ['zh-CN', 'zh', 'en']});
                window.chrome = {runtime: {}};
            """)
            await context.set_extra_http_headers({
                "Referer": "https://hrewards.huazhu.com/",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            })
            page = await context.new_page()

            logger.info(f"访问详情页: {detail_url}")
            try:
                await page.goto(detail_url, wait_until="domcontentloaded", timeout=30000)
            except Exception as e:
                logger.warning(f"页面加载异常: {e}")
                await browser.close()
                return {"error": f"页面加载超时: {e}"}

            # 等待目标元素出现，而非固定等待
            try:
                await page.wait_for_selector(".booking_item", timeout=15000)
            except Exception:
                # 回退到通用选择器
                for sel in ["[class*='booking'] li", "[class*='booking'] > div", "[class*='room']"]:
                    try:
                        await page.wait_for_selector(sel, timeout=5000)
                    except Exception:
                        continue

            rooms = []
            booking_items = await page.query_selector_all(".booking_item")

            if not booking_items:
                for sel in ["[class*='booking'] li", "[class*='booking'] > div", "[class*='room']"]:
                    booking_items = await page.query_selector_all(sel)
                    if booking_items:
                        break

            for item in booking_items:
                try:
                    room_name_el = await item.query_selector(".room_name")
                    room_name = (await room_name_el.text_content()).strip() if room_name_el else ""
                    price_el = await item.query_selector(".original_price")
                    price = (await price_el.text_content()).strip() if price_el else ""
                    if room_name:
                        rooms.append({"房型名": room_name, "价格": price})
                except Exception:
                    continue

            await browser.close()
            return rooms

    try:
        rooms = _run_async(_crawl())
    except Exception as e:
        return {"error": f"爬取失败: {e}"}

    if not rooms or isinstance(rooms, dict) and "error" in rooms:
        return {
            "酒店名": matched_name,
            "酒店ID": hotel_id,
            "城市": matched["城市名"],
            "品牌": matched["品牌名"],
            "房型": [],
            "error": "未获取到房型信息" if not isinstance(rooms, dict) else rooms.get("error", ""),
        }

    return {
        "酒店名": matched_name,
        "酒店ID": hotel_id,
        "城市": matched["城市名"],
        "品牌": matched["品牌名"],
        "房型": rooms,
    }


def _parse_price(price_str: str):
    """从价格字符串中提取数字，失败返回 None"""
    try:
        return float(re.sub(r'[^\d.]', '', price_str))
    except (ValueError, TypeError):
        return None


def query_hotel(days_in: int, days_out: int, hotel_name: str):
    """执行一次查询"""
    date_in = datetime.now() + timedelta(days=days_in)
    date_out = datetime.now() + timedelta(days=days_out)

    if date_in >= date_out:
        print("  错误: 离店日期必须晚于入住日期")
        return

    nights = (date_out - date_in).days
    print(f"酒店: {hotel_name}")
    print(f"入住: {date_in.strftime('%Y-%m-%d')} (距今{days_in}天)")
    print(f"离店: {date_out.strftime('%Y-%m-%d')} (距今{days_out}天)")
    print(f"共 {nights} 晚")
    print("=" * 60)

    # 逐晚查询，收集数据
    night_dates = []
    night_results = []
    current = date_in
    idx = 1
    while current < date_out:
        next_day = current + timedelta(days=1)
        ci = current.strftime("%Y-%m-%d")
        co = next_day.strftime("%Y-%m-%d")
        night_dates.append(ci)

        try:
            result = _query_night_price_sync(hotel_name, ci, co)
            night_results.append(result)
        except Exception as e:
            night_results.append({"error": str(e)})

        current = next_day
        idx += 1

    # 逐晚打印详情
    for i, result in enumerate(night_results):
        ci = night_dates[i]
        print(f"\n【第{i+1}晚】{ci}")
        print("-" * 40)

        if "error" in result and "房型" not in result:
            print(f"  查询失败: {result['error']}")
            continue

        print(f"  酒店: {result.get('酒店名', hotel_name)}")
        print(f"  品牌: {result.get('品牌', '-')}")
        print(f"  城市: {result.get('城市', '-')}")
        rooms = result.get("房型", [])
        if rooms:
            for room in rooms:
                raw_price = room['价格']
                price_num = _parse_price(raw_price)
                if price_num is not None:
                    discount_price = price_num * 0.75
                    print(f"  {room['房型名']}: ¥{price_num:.0f}  75折: ¥{discount_price:.0f}")
                else:
                    print(f"  {room['房型名']}: ¥{raw_price}")
        else:
            print("  无可用房型")
            if result.get("error"):
                print(f"  ({result['error']})")

    # 汇总表格：房型 × 日期，75折价格
    print("\n" + "=" * 80)
    print("汇总（75折价格）")
    print("=" * 80)

    # 提取所有房型名称（取第一晚成功的结果）
    all_rooms = []
    for result in night_results:
        rooms = result.get("房型", [])
        if rooms:
            all_rooms = rooms
            break

    if not all_rooms:
        print("无可用房型数据")
        print("查询完成\n")
        return

    # 构建 房型 → [每晚75折价格或 None]
    room_names = [r['房型名'] for r in all_rooms]
    room_prices = {r['房型名']: [] for r in all_rooms}

    for result in night_results:
        rooms = result.get("房型", [])
        room_price_map = {}
        for r in rooms:
            price_num = _parse_price(r['价格'])
            if price_num is not None:
                room_price_map[r['房型名']] = price_num * 0.75
        for name in room_names:
            room_prices[name].append(room_price_map.get(name, None))

    def _vis_len(s):
        """终端视觉宽度：CJK字符=2，其余=1"""
        return sum(2 if '一' <= ch <= '鿿' else 1 for ch in s)

    def _pad(s, target_w):
        """填充空格到指定视觉宽度"""
        cur = _vis_len(s)
        if cur >= target_w:
            return s
        return s + " " * (target_w - cur)

    date_headers = [d[5:] for d in night_dates]  # 去掉年份，只保留 MM-DD
    price_col_w = 7
    room_col_target = max(_vis_len("房型"), max(_vis_len(n) for n in room_names)) + 2

    # 表头
    header = _pad("房型", room_col_target)
    for d in date_headers:
        header += f"{d:>{price_col_w}}"
    header += " " + _pad("总计", price_col_w)
    print(header)
    print("-" * _vis_len(header))

    # 每行
    for name in room_names:
        row_str = _pad(name, room_col_target)
        total = 0.0
        for p in room_prices[name]:
            if p is not None:
                row_str += f"{p:>{price_col_w}.0f}"
                total += p
            else:
                row_str += f"{'—':>{price_col_w}}"
        row_str += f"{total:>{price_col_w}.0f}"
        print(row_str)

    # 合计行
    total_row = _pad("合计", room_col_target)
    grand_total = 0.0
    for d_idx in range(len(night_dates)):
        night_total = 0.0
        night_valid = True
        for name in room_names:
            p = room_prices[name][d_idx]
            if p is None:
                night_valid = False
                break
            night_total += p
        if night_valid:
            total_row += f"{night_total:>{price_col_w}.0f}"
            grand_total += night_total
        else:
            total_row += f"{'—':>{price_col_w}}"
    total_row += f"{grand_total:>{price_col_w}.0f}"
    print(total_row)
    print("=" * 80)
    print("查询完成\n")


def main():
    print("=" * 60)
    print("  酒店房价查询工具")
    print("=" * 60)
    print("输入格式: 入住天数 离店天数 酒店名")
    print("示例: 10 12 全季酒店")
    print("退出: quit / exit")
    print("=" * 60)

    # 支持命令行一次性运行
    if len(sys.argv) >= 4:
        query_hotel(int(sys.argv[1]), int(sys.argv[2]), " ".join(sys.argv[3:]))

    while True:
        try:
            line = input("\n> ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n再见!")
            break

        if not line:
            continue
        if line.lower() in ("quit", "exit", "q"):
            print("再见!")
            break

        parts = line.split()
        if len(parts) < 3:
            print("  输入格式: 入住天数 离店天数 酒店名")
            print("  示例: 10 12 全季酒店")
            continue

        try:
            days_in = int(parts[0])
            days_out = int(parts[1])
        except ValueError:
            print("  前两个参数必须是数字")
            continue

        hotel_name = " ".join(parts[2:])
        query_hotel(days_in, days_out, hotel_name)


if __name__ == "__main__":
    main()
