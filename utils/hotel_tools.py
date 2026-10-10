"""
酒店数据工具 — 华住会酒店查询与爬取
"""
import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import difflib
import json
import os
import re
from playwright_stealth import Stealth
from typing import Any, Dict, List, Optional


from loguru import logger
from openpyxl import Workbook, load_workbook


def _run_async(coro):
    """在独立线程中运行 async 协程，避免与主 event loop 冲突"""
    with ThreadPoolExecutor() as pool:
        return pool.submit(asyncio.run, coro).result()


DEFAULT_EXCEL_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "huazhu_hotels.xlsx")
HOTELS_DB_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "hotels.xlsx")

# 华住会品牌词典（按长度降序，优先匹配最长品牌）
HUAZHU_BRANDS = sorted([
    "桔子水晶", "桔子精选", "桔子酒店",
    "全季", "全季大观",
    "汉庭优佳", "汉庭",
    "漫心", "漫心府",
    "美居", "美居酒店",
    "宜必思尚品", "宜必思",
    "星程", "星程酒店",
    "CitiGO", "Citigo", "花间堂", "花间",
    "禧玥", "施柏阁", "城家",
], key=len, reverse=True)


def _normalize(name: str) -> str:
    """清洗酒店名：去括号、去符号、中文数字转阿拉伯数字、去冗余字"""
    name = re.sub(r'\([^)]*\)|（[^）]*）|\[[^\]]*\]|【[^】]*】', '', name)
    name = re.sub(r'[^一-龥a-zA-Z0-9]', '', name)
    cn_num_map = {'零': 0, '一': 1, '二': 2, '三': 3, '四': 4,
                  '五': 5, '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}

    def replace_cn(match):
        cn = match.group(0)
        if len(cn) == 2 and cn[0] in cn_num_map and cn[1] == '十':
            return str(cn_num_map[cn[0]] * 10)
        elif len(cn) == 2 and cn[0] == '十' and cn[1] in cn_num_map:
            return str(10 + cn_num_map[cn[1]])
        elif len(cn) == 3 and cn[1] == '十':
            return str(cn_num_map[cn[0]] * 10 + cn_num_map[cn[2]])
        return str(cn_num_map.get(cn[0], 0))

    name = re.sub(r'[零一二三四五六七八九十]{1,3}', replace_cn, name)
    for word in ['酒店', '分店', '大厦店', '楼店', '层店', '新店', '店']:
        name = name.replace(word, '')
    return name.strip()


def _read_excel(path: str) -> List[Dict[str, Any]]:
    """读取华住会 Excel 文件"""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    if ws.max_row <= 1:
        wb.close()
        return []
    hotels = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        hotels.append({
            "酒店名": row[0] or "",
            "门市价格": row[1] or "",
            "入住日期": row[2] or "",
            "离店日期": row[3] or "",
            "酒店地址": row[4] or "",
            "酒店类型": row[5] or "",
        })
    wb.close()
    return hotels


def _find_best_match(hotels: List[Dict], keyword: str, threshold: float = 0.4) -> List[Dict]:
    """使用 difflib 模糊匹配，返回最接近的酒店列表"""
    kw = _normalize(keyword)
    scored = []
    for h in hotels:
        norm = _normalize(h["酒店名"])
        ratio = difflib.SequenceMatcher(None, kw, norm).ratio()
        if kw in norm or norm in kw:
            ratio = max(ratio, 0.85)
        if ratio >= threshold:
            scored.append((ratio, h))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [{"score": round(s, 2), **h} for s, h in scored[:5]]


# ---------- Tool 1: 搜索酒店 + 实时查价（合并） ----------

def search_hotel_price(hotel_name: str, check_in: str, check_out: str) -> str:
    """
    根据酒店名查询ID，爬取华住会详情页获取所有房型及实时价格
    """
    from playwright.async_api import async_playwright

    # 1. 从 hotels.xlsx 查找酒店ID
    hotel_db = _load_hotel_db()
    if not hotel_db:
        return json.dumps({"error": "酒店ID数据库为空，请确认 data/hotels.xlsx 存在"}, ensure_ascii=False)

    matched = _find_hotel_by_name(hotel_db, hotel_name)
    if not matched:
        return json.dumps({"error": f"未找到 '{hotel_name}' 对应的酒店记录"}, ensure_ascii=False)

    hotel_id = matched["酒店ID"]
    matched_name = matched["酒店名"]
    logger.debug(f"[search_hotel_price] 匹配酒店: {matched_name}, ID: {hotel_id}")

    # 2. 拼接详情页URL
    detail_url = (
        f"https://hrewards.huazhu.com/hotel/detail?"
        f"checkInDate={check_in}&checkOutDate={check_out}&hotelId={hotel_id}"
    )

    async def _crawl():
        async with async_playwright() as p:
            browser = await p.chromium.launch(
                headless=True,
                args=[
                    "--disable-blink-features=AutomationControlled",
                    "--disable-dev-shm-usage",
                    "--no-sandbox",
                    "--disable-gpu",
                    "--disable-setuid-sandbox",
                    "--disable-web-security",
                ]
            )
            context = await browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                viewport={"width": 1920, "height": 1080},
                locale="zh-CN",
                timezone_id="Asia/Shanghai",
            )
            await context.add_init_script("""
                Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
                Object.defineProperty(navigator, 'plugins', {get: () => [1, 2, 3, 4, 5]});
                Object.defineProperty(navigator, 'languages', {get: () => ['zh-CN', 'zh', 'en']});
                window.chrome = {runtime: {}};
            """)
            await context.set_extra_http_headers({
                "Referer": "https://hrewards.huazhu.com/",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            })
            page = await context.new_page()
            await Stealth().apply_stealth_async(page)

            logger.debug(f"[search_hotel_price] 访问详情页: {detail_url}")
            debug_dir = os.path.join(os.path.dirname(__file__), "..", "data")
            try:
                await page.goto(detail_url, wait_until="domcontentloaded", timeout=30000)
            except Exception as e:
                logger.warning(f"[search_hotel_price] 页面加载异常: {e}")
                await browser.close()
                return {"error": f"页面加载超时: {e}"}

            await asyncio.sleep(8)

            # 滚动触发懒加载
            await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            await asyncio.sleep(3)
            await page.evaluate("window.scrollTo(0, 0)")
            await asyncio.sleep(2)

            # 提取房型信息
            rooms = []
            booking_items = await page.query_selector_all(".booking_item")
            logger.debug(f"[search_hotel_price] 找到 {len(booking_items)} 个 booking_item")

            if not booking_items:
                for sel in ["[class*='booking'] li", "[class*='booking'] > div", "[class*='room']"]:
                    booking_items = await page.query_selector_all(sel)
                    if booking_items:
                        break

            for item in booking_items:
                try:
                    room_name_el = await item.query_selector(".room_name")
                    room_name = (await room_name_el.text_content()).strip() if room_name_el else ""
                    price_el = await item.query_selector(".original_price")
                    price = (await price_el.text_content()).strip() if price_el else ""
                    if room_name:
                        rooms.append({"房型名": room_name, "价格": price})
                except Exception:
                    continue

            await browser.close()
            return rooms

    try:
        rooms = _run_async(_crawl())
    except Exception as e:
        return json.dumps({"error": f"爬取失败: {e}"}, ensure_ascii=False)

    if not rooms:
        return json.dumps({
            "酒店名": matched_name,
            "酒店ID": hotel_id,
            "城市": matched["城市名"],
            "品牌": matched["品牌名"],
            "error": "未获取到房型信息，可能满房或页面结构变化",
        }, ensure_ascii=False)

    # 保存到本地 Excel
    _save_price_to_excel(matched_name, hotel_id, check_in, check_out, rooms)

    return json.dumps({
        "酒店名": matched_name,
        "酒店ID": hotel_id,
        "城市": matched["城市名"],
        "品牌": matched["品牌名"],
        "入住日期": check_in,
        "离店日期": check_out,
        "房型": rooms,
    }, ensure_ascii=False, indent=2)


# ---------- 辅助函数：酒店ID数据库查找 ----------

def _load_hotel_db(path: str = HOTELS_DB_PATH) -> List[Dict[str, str]]:
    """加载 hotels.xlsx 酒店ID数据库，返回 [{酒店名, 酒店ID, 城市名, 城市ID, 品牌名, 品牌类型}]"""
    if not os.path.exists(path):
        logger.warning(f"[search_hotel_price] 酒店ID数据库不存在: {path}")
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    if ws.max_row <= 1:
        wb.close()
        return []
    hotels = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        hotels.append({
            "酒店名": str(row[0] or ""),
            "酒店ID": str(row[1] or ""),
            "城市名": str(row[2] or ""),
            "城市ID": str(row[3] or ""),
            "品牌名": str(row[4] or ""),
            "品牌类型": str(row[5] or ""),
        })
    wb.close()
    return hotels


def _find_hotel_by_name(hotels: List[Dict], keyword: str, threshold: float = 0.5) -> Optional[Dict]:
    """从酒店ID数据库中模糊匹配酒店，返回最佳匹配"""
    kw = _normalize(keyword)
    best = None
    best_score = 0.0
    for h in hotels:
        name = _normalize(h["酒店名"])
        if kw in name or name in kw:
            score = 0.9
        else:
            score = difflib.SequenceMatcher(None, kw, name).ratio()
        if score > best_score and score >= threshold:
            best_score = score
            best = h
    return best


def _save_price_to_excel(hotel_name: str, hotel_id: str, check_in: str, check_out: str, rooms: list,
                         path: str = None) -> None:
    """保存查价结果到 Excel，同酒店同日期覆盖旧数据，不同追加"""
    if path is None:
        path = os.path.join(os.path.dirname(__file__), "..", "data", "hotel_prices.xlsx")

    headers = ["酒店名", "酒店ID", "入住日期", "离店日期", "房型名", "价格", "查询时间"]
    query_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # 读取已有数据
    existing_rows = []
    if os.path.exists(path):
        try:
            wb = load_workbook(path)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    existing_rows.append(list(row))
            wb.close()
        except Exception:
            existing_rows = []

    # 移除同酒店同日期的旧记录
    filtered = [r for r in existing_rows if not (r[0] == hotel_name and r[2] == check_in)]

    # 追加新记录
    for room in rooms:
        filtered.append([hotel_name, hotel_id, check_in, check_out, room["房型名"], room["价格"], query_time])

    # 写入 Excel
    dir_path = os.path.dirname(path)
    if dir_path and not os.path.exists(dir_path):
        os.makedirs(dir_path, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "酒店查价记录"
    ws.append(headers)
    for row in filtered:
        ws.append(row)
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 20
    wb.save(path)
    logger.debug(f"[get_hotel_price] 查价结果已保存至 {path}（{len(rooms)} 条房型）")


# ---------- Tool 3: 批量爬取刷新 ----------

def crawl_hotels(city: str, check_in: str, check_out: str,
                 excel_path: str = DEFAULT_EXCEL_PATH,
                 page_size: int = 100) -> str:
    """
    批量爬取华住会指定城市全部酒店数据，保存至本地 Excel
    """
    from playwright.async_api import async_playwright

    async def _crawl() -> List[Dict]:
        all_hotels: List[Dict] = []
        current_page = 1
        max_pages = 100

        url = (f"https://hrewards.huazhu.com/hotel?cityName={city}&"
               f"checkInDate={check_in}&checkOutDate={check_out}&"
               f"keyword=&pageIndex=1&pageSize={page_size}")

        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True, args=[
                "--disable-blink-features=AutomationControlled", "--disable-dev-shm-usage", "--no-sandbox",
            ])
            context = await browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
                viewport={"width": 1920, "height": 1080}, locale="zh-CN",
            )
            await context.set_extra_http_headers({
                "Referer": "https://hrewards.huazhu.com/",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            })
            page = await context.new_page()
            await Stealth().apply_stealth_async(page)
            await page.goto(url, wait_until="domcontentloaded", timeout=30000)

            while current_page <= max_pages:
                try:
                    await page.wait_for_selector(
                        ".hotel-item, .hotel-list-item, .hotel-card, [class*='hotel']",
                        timeout=15000,
                    )
                except Exception:
                    await page.wait_for_load_state("networkidle", timeout=10000)

                await page.wait_for_timeout(2000)

                page_hotels = []
                try:
                    li_els = await page.query_selector_all(
                        "xpath=/html/body/div/div/div[1]/div/div/div/div[2]/div[2]/div/ul/li"
                    )
                    for li in li_els:
                        try:
                            name_el = await li.query_selector(
                                ".hotel-name, .name, .title, [class*='name']"
                            )
                            hotel_name = ""
                            if name_el:
                                hotel_name = (await name_el.text_content()).strip()

                            price_el = await li.query_selector(
                                "xpath=./div[2]/div[2]/div[1]/div"
                            )
                            price = ""
                            if price_el:
                                t = (await price_el.text_content()).strip()
                                price = t if t else ""

                            addr_el = await li.query_selector(
                                ".hotel-address, .address, .location, [class*='address']"
                            )
                            addr = ""
                            if addr_el:
                                addr = (await addr_el.text_content()).strip()

                            type_el = await li.query_selector(
                                ".hotel-type, .brand, .brand-name, [class*='brand']"
                            )
                            hotel_type = ""
                            if type_el:
                                hotel_type = (await type_el.text_content()).strip()

                            if hotel_name:
                                page_hotels.append({
                                    "酒店名": hotel_name,
                                    "门市价格": price,
                                    "入住日期": check_in,
                                    "离店日期": check_out,
                                    "酒店地址": addr,
                                    "酒店类型": hotel_type,
                                })
                        except Exception:
                            continue
                except Exception:
                    pass

                if not page_hotels:
                    break

                sold_out = sum(1 for h in page_hotels if not h.get("门市价格"))
                ratio = sold_out / len(page_hotels) if page_hotels else 0
                all_hotels.extend(page_hotels)

                if ratio > 0.8:
                    break

                try:
                    next_btn = await page.query_selector(
                        "xpath=//button[contains(text(), '下一页') or contains(text(), '>')]"
                    )
                    if not next_btn:
                        break
                    is_disabled = await next_btn.get_attribute("disabled")
                    if is_disabled:
                        break
                    await next_btn.scroll_into_view_if_needed()
                    await page.wait_for_timeout(500)
                    await next_btn.click()
                    await page.wait_for_timeout(4000)
                    current_page += 1
                except Exception:
                    break

            await browser.close()
            return all_hotels

    hotels = _run_async(_crawl())

    if not hotels:
        return json.dumps({"error": f"爬取 {city} 酒店数据失败"}, ensure_ascii=False)

    _save_to_excel(hotels, excel_path)
    return json.dumps({
        "message": f"成功爬取 {len(hotels)} 家酒店",
        "城市": city, "入住": check_in, "离店": check_out,
        "保存路径": os.path.abspath(excel_path),
    }, ensure_ascii=False)


def _save_to_excel(hotels: List[Dict], path: str) -> None:
    """保存酒店列表到 Excel"""
    dir_path = os.path.dirname(path)
    if dir_path and not os.path.exists(dir_path):
        os.makedirs(dir_path, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "华住会酒店"
    headers = ["酒店名", "门市价格", "入住日期", "离店日期", "酒店地址", "酒店类型"]
    ws.append(headers)
    for h in hotels:
        ws.append([h.get(k, "") for k in headers])

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 20
    wb.save(path)
    logger.info(f"华住会酒店数据已保存至 {path}（{len(hotels)} 条）")
